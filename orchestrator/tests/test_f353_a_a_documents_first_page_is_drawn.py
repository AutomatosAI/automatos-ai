"""F353 (issue #947): a document's first page is drawn as a small PNG.

Every file here is made in the test: a PDF (WeasyPrint), a Word document
(python-docx), a sheet (openpyxl), a CSV and markdown. Each comes out as a PNG
``THUMBNAIL_WIDTH_PX`` wide with ink on it; a very tall page is cut short; the
child-process render gives the same picture, and a broken file fails with its
reason instead of a picture. A file's image URLs are never fetched.
"""
from __future__ import annotations

import io

import pytest
from PIL import Image

from modules.documents.thumbnails import render
from modules.documents.thumbnails.html_sources import document_html
from modules.documents.thumbnails.render import (
    MAX_THUMBNAIL_HEIGHT_PX,
    THUMBNAIL_WIDTH_PX,
    ThumbnailError,
    render_png,
    render_png_isolated,
)


def _pdf(html: str = "<h1>Harbourline invoice</h1><p>Two crates of oranges</p>") -> bytes:
    from weasyprint import HTML

    return HTML(string=html).write_pdf()


def _image(png: bytes) -> Image.Image:
    assert png.startswith(b"\x89PNG"), "not a PNG"
    return Image.open(io.BytesIO(png))


def _has_ink(img: Image.Image) -> bool:
    darkest, _ = img.convert("L").getextrema()
    return darkest < 128


def _docx() -> bytes:
    from docx import Document

    doc = Document()
    doc.add_heading("Quote for Harbourline", 0)
    doc.add_paragraph("Delivery on Friday.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Item", "Price"
    table.cell(1, 0).text, table.cell(1, 1).text = "Oranges", "12.00"
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _xlsx() -> bytes:
    import openpyxl

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.append(["Item", "Qty", "Price"])
    sheet.append(["Oranges", 2, 12.0])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_a_pdfs_first_page_is_a_png_of_the_card_width_with_ink_on_it():
    img = _image(render_png(_pdf(), ".pdf"))
    assert img.width == THUMBNAIL_WIDTH_PX
    assert img.height > img.width  # an A4 page, portrait
    assert _has_ink(img)


def test_a_very_tall_page_is_cut_to_the_top():
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=595, height=6000)
    out = io.BytesIO()
    writer.write(out)
    img = _image(render_png(out.getvalue(), ".pdf"))
    assert img.width == THUMBNAIL_WIDTH_PX
    assert img.height <= MAX_THUMBNAIL_HEIGHT_PX + 1


@pytest.mark.parametrize(
    "ext, data",
    [
        (".docx", _docx()),
        (".xlsx", _xlsx()),
        (".csv", b"Item,Qty\nOranges,2\nLemons,5\n"),
        (".md", b"# Weekly report\n\n- 12 orders shipped\n- 2 returns\n"),
        (".txt", b"Plain notes for the week."),
    ],
)
def test_a_word_document_a_sheet_a_csv_and_markdown_are_drawn(ext, data):
    img = _image(render_png(data, ext))
    assert img.width == THUMBNAIL_WIDTH_PX
    assert _has_ink(img)


def test_a_sheets_values_are_escaped_into_the_page():
    page = document_html(b"Name\n<script>alert(1)</script>\n", ".csv")
    assert "<script>" not in page
    assert "&lt;script&gt;" in page


def test_a_documents_image_urls_are_never_fetched(monkeypatch):
    asked = []

    def refuse(url, *args, **kwargs):
        asked.append(url)
        raise ValueError("refused")

    monkeypatch.setattr(render, "_refuse_fetch", refuse)
    markdown = b"# Report\n\n![x](http://169.254.169.254/latest/meta-data/)\n"
    assert _image(render_png(markdown, ".md")).width == THUMBNAIL_WIDTH_PX
    assert asked, "the image URL never reached the refusing fetcher"
    assert all(u.startswith("http://169.254.169.254") for u in asked)


def test_an_unknown_type_or_an_empty_file_has_no_picture():
    with pytest.raises(ThumbnailError):
        render_png(b"PK\x03\x04", ".pptx")
    with pytest.raises(ThumbnailError):
        render_png(b"", ".pdf")


def test_the_child_process_render_draws_the_same_page():
    img = _image(render_png_isolated(_pdf(), ".pdf"))
    assert img.width == THUMBNAIL_WIDTH_PX
    assert _has_ink(img)


def test_a_broken_file_fails_with_its_reason_not_a_picture():
    with pytest.raises(ThumbnailError) as caught:
        render_png_isolated(b"this is not a PDF", ".pdf")
    assert "not a readable PDF" in str(caught.value)
