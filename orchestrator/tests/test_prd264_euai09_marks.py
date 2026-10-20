"""PRD-264 EUAI-09 — documents and images carry a machine-readable AI mark (Art. 50(2)).

``modules/documents/provenance_marks.py`` writes the IPTC Digital Source Type where each
format keeps its metadata, and every mark carries a random per-file instance id (FR-16).
Each format is made by the library the platform makes it with, marked, and read back:

* **PDF** (pypdf, and a real WeasyPrint page): an XMP packet in the catalogue's
  ``/Metadata`` stream, ``Iptc4xmpExt:DigitalSourceType`` and ``xmpMM:InstanceID``; the
  pages and the ``/Info`` author are kept.
* **DOCX** (python-docx) and **XLSX** (``write_xlsx``, XlsxWriter): the custom
  properties ``AIGenerated`` (the IPTC term) and ``AIInstanceID`` in
  ``docProps/custom.xml``, named in ``[Content_Types].xml`` and ``_rels/.rels``; the
  package still opens and the workspace's own custom properties are kept.
* **PNG** (an ``iTXt`` chunk) and **JPEG** (an ``APP1`` segment), both made by Pillow:
  the pixels decode the same; an image with a C2PA manifest is kept byte for byte; an
  XMP packet that already states a source type keeps it.

The mark holds the term and the random id only (Guidelines para. 94: a packet's fields
are checked against an allowlist), and two marks of identical input never give the
same SHA-256 (FR-16), so a predictable document cannot be confirmed by guessing hashes.
"""
from __future__ import annotations

import hashlib
import io
import shutil
import struct
import uuid
import zipfile
import zlib
from pathlib import Path

import pytest
from lxml import etree

from modules.documents import provenance_marks as pm

COMPOSITE = "http://cv.iptc.org/newscodes/digitalsourcetype/compositeWithTrainedAlgorithmicMedia"
TRAINED = "http://cv.iptc.org/newscodes/digitalsourcetype/trainedAlgorithmicMedia"
RDF = "http://www.w3.org/1999/02/22-rdf-syntax-ns#"
DST_TAG = "{http://iptc.org/std/Iptc4xmpExt/2008-02-29/}DigitalSourceType"
INSTANCE_TAG = "{http://ns.adobe.com/xap/1.0/mm/}InstanceID"
MARK_FIELDS = {DST_TAG, INSTANCE_TAG}
CUSTOM = "{http://schemas.openxmlformats.org/officeDocument/2006/custom-properties}"
VT = "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}"
AUTHOR = "Harbourline Coffee Roasters"
SHEET = {"columns": ["Coffee", "Kg left"], "rows": [["Kirinyaga", 41], ["Guji Shakiso", 55.5]]}


# --------------------------------------------------------------------------- readers


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _fields(packet: bytes) -> dict:
    """{property tag: text} over every rdf:Description in an XMP packet."""
    root = etree.fromstring(packet)
    found = {}
    for description in root.iter(f"{{{RDF}}}Description"):
        for prop in description:
            found[prop.tag] = prop.text
    return found


def _assert_instance_id(value: str) -> None:
    assert value.startswith("xmp.iid:")
    assert uuid.UUID(value[len("xmp.iid:"):]).version == 4


def _assert_mark(packet: bytes, term: str) -> None:
    fields = _fields(packet)
    assert set(fields) == MARK_FIELDS  # the allowlist: nothing about a user, workspace or prompt
    assert fields[DST_TAG] == term
    _assert_instance_id(fields[INSTANCE_TAG])


def _pdf_packet(path: Path) -> bytes:
    from pypdf import PdfReader

    return PdfReader(str(path)).trailer["/Root"]["/Metadata"].get_object().get_data()


def _custom_properties(path: Path) -> dict:
    with zipfile.ZipFile(path) as package:
        root = etree.fromstring(package.read("docProps/custom.xml"))
    return {prop.get("name"): prop.find(f"{VT}lpwstr").text for prop in root.iter(f"{CUSTOM}property")}


def _png_chunks(data: bytes) -> list:
    """(type, body) per chunk, each CRC checked: an independent reader, not the marker's."""
    chunks, at = [], 8
    while at < len(data):
        length, kind = struct.unpack(">I4s", data[at:at + 8])
        body = data[at + 8:at + 8 + length]
        (crc,) = struct.unpack(">I", data[at + 8 + length:at + 12 + length])
        assert crc == zlib.crc32(kind + body) & 0xFFFFFFFF
        chunks.append((kind, body))
        at += 12 + length
    return chunks


def _png_packet(data: bytes) -> bytes:
    (body,) = [b for kind, b in _png_chunks(data) if kind == b"iTXt" and b.startswith(b"XML:com.adobe.xmp\x00")]
    return body[len(b"XML:com.adobe.xmp\x00\x00\x00\x00\x00"):]


def _jpeg_segments(data: bytes) -> list:
    """(marker, body) per segment before the first scan: an independent reader."""
    assert data.startswith(b"\xff\xd8")
    segments, at = [], 2
    while data[at + 1] != 0xDA:
        (length,) = struct.unpack(">H", data[at + 2:at + 4])
        segments.append((data[at + 1], data[at + 4:at + 2 + length]))
        at += 2 + length
    return segments


def _jpeg_packet(data: bytes) -> bytes:
    header = b"http://ns.adobe.com/xap/1.0/\x00"
    (body,) = [b for marker, b in _jpeg_segments(data) if marker == 0xE1 and b.startswith(header)]
    return body[len(header):]


# --------------------------------------------------------------------------- fixtures


def _pdf(path: Path) -> Path:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    writer.add_metadata({"/Author": AUTHOR, "/Title": "Invoice HL-1"})
    with open(path, "wb") as out:
        writer.write(out)
    return path


def _docx(path: Path) -> Path:
    import docx

    document = docx.Document()
    document.add_paragraph("Dear Ana, your order has shipped.")
    document.save(str(path))
    return path


def _xlsx(path: Path) -> Path:
    from modules.documents.xlsx_render import write_xlsx

    write_xlsx(str(path), "Green coffee stock", SHEET, None)
    return path


def _image(fmt: str) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (8, 8), "#1E3A5F").save(buf, fmt)
    return buf.getvalue()


def _pixels(data: bytes) -> bytes:
    from PIL import Image

    with Image.open(io.BytesIO(data)) as picture:
        return picture.convert("RGB").tobytes()


def _png_with(chunk_type: bytes, body: bytes) -> bytes:
    """The Pillow PNG with one more chunk straight after IHDR."""
    data = _image("PNG")
    after_ihdr = 8 + 12 + struct.unpack(">I", data[8:12])[0]
    chunk = struct.pack(">I", len(body)) + chunk_type + body + struct.pack(">I", zlib.crc32(chunk_type + body) & 0xFFFFFFFF)
    return data[:after_ihdr] + chunk + data[after_ihdr:]


def _jpeg_with(marker: int, body: bytes) -> bytes:
    """The Pillow JPEG with one more segment straight after SOI."""
    data = _image("JPEG")
    return data[:2] + bytes((0xFF, marker)) + struct.pack(">H", len(body) + 2) + body + data[2:]


# --------------------------------------------------------------------------- the packet


def test_a_new_packet_holds_the_term_and_a_random_instance_id_only():
    packet = pm.xmp_packet(TRAINED)
    _assert_mark(packet, TRAINED)
    assert _fields(pm.xmp_packet(TRAINED))[INSTANCE_TAG] != _fields(packet)[INSTANCE_TAG]


def test_only_the_iptc_ai_terms_are_written():
    with pytest.raises(ValueError):
        pm.xmp_packet("http://cv.iptc.org/newscodes/digitalsourcetype/digitalCapture")


def test_a_packet_that_states_its_source_keeps_it_and_gains_an_instance_id():
    theirs = pm.xmp_packet(TRAINED).replace(TRAINED.encode(), b"http://example.test/their-own-term")
    without_id = theirs.replace(b"<xmpMM:InstanceID>", b"<xmpMM:Other>").replace(b"</xmpMM:InstanceID>", b"</xmpMM:Other>")
    kept = pm.xmp_packet(COMPOSITE, without_id)
    fields = _fields(kept)
    assert fields[DST_TAG] == "http://example.test/their-own-term"
    _assert_instance_id(fields[INSTANCE_TAG])
    assert pm.xmp_packet(COMPOSITE, theirs) is None  # it states both: nothing to add


# --------------------------------------------------------------------------- PDF


def test_a_pdf_carries_the_mark_in_its_xmp_and_keeps_its_pages_and_author(tmp_path):
    from pypdf import PdfReader

    path = _pdf(tmp_path / "invoice.pdf")
    sha = pm.mark_document(str(path), "pdf", COMPOSITE)

    _assert_mark(_pdf_packet(path), COMPOSITE)
    reader = PdfReader(str(path))
    assert len(reader.pages) == 1 and reader.metadata.author == AUTHOR
    assert sha == _sha(path)


def test_a_weasyprint_pdf_is_marked_and_still_reads(tmp_path):
    from pypdf import PdfReader
    from weasyprint import HTML

    path = tmp_path / "letter.pdf"
    HTML(string="<h1>Invoice HL-0142</h1><p>Thank you for your order.</p>").write_pdf(str(path))
    pm.mark_document(str(path), "pdf", COMPOSITE)

    _assert_mark(_pdf_packet(path), COMPOSITE)
    assert "Invoice HL-0142" in PdfReader(str(path)).pages[0].extract_text()


# --------------------------------------------------------------------------- DOCX and XLSX


def _assert_ooxml_mark(path: Path) -> dict:
    props = _custom_properties(path)
    assert props["AIGenerated"] == COMPOSITE
    _assert_instance_id(props["AIInstanceID"])
    with zipfile.ZipFile(path) as package:
        assert b'PartName="/docProps/custom.xml"' in package.read("[Content_Types].xml")
        assert b'Target="docProps/custom.xml"' in package.read("_rels/.rels") or \
            b'Target="/docProps/custom.xml"' in package.read("_rels/.rels")
    return props


def test_a_docx_carries_the_custom_property_and_still_opens(tmp_path):
    import docx

    path = _docx(tmp_path / "letter.docx")
    sha = pm.mark_document(str(path), "docx", COMPOSITE)

    _assert_ooxml_mark(path)
    assert "Dear Ana, your order has shipped." in [p.text for p in docx.Document(str(path)).paragraphs]
    assert sha == _sha(path)


def test_an_xlsx_carries_the_custom_property_and_still_opens(tmp_path):
    import openpyxl

    path = _xlsx(tmp_path / "stock.xlsx")
    pm.mark_document(str(path), "xlsx", COMPOSITE)

    _assert_ooxml_mark(path)
    values = [cell.value for row in openpyxl.load_workbook(str(path)).active.iter_rows() for cell in row]
    assert "Kirinyaga" in values and 55.5 in values


def _with_their_custom_property(path: Path) -> Path:
    """The package with a custom property of its own (pid 2), as an uploaded template has."""
    custom = (b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
              b'<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/custom-properties" '
              b'xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes">'
              b'<property fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}" pid="2" name="Department">'
              b'<vt:lpwstr>Roastery</vt:lpwstr></property></Properties>')
    with zipfile.ZipFile(path) as source:
        parts = {info.filename: source.read(info) for info in source.infolist()}
    parts["docProps/custom.xml"] = custom
    out = path.with_name("theirs-" + path.name)
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as target:
        for name, body in parts.items():
            target.writestr(name, body)
    return out


def test_a_package_keeps_its_own_custom_properties_and_a_second_mark_replaces_the_first(tmp_path):
    path = _with_their_custom_property(_docx(tmp_path / "template.docx"))
    pm.mark_document(str(path), "docx", COMPOSITE)
    first = _assert_ooxml_mark(path)
    pm.mark_document(str(path), "docx", COMPOSITE)
    second = _assert_ooxml_mark(path)

    assert second["Department"] == "Roastery"
    assert set(second) == {"Department", "AIGenerated", "AIInstanceID"}
    assert second["AIInstanceID"] != first["AIInstanceID"]
    with zipfile.ZipFile(path) as package:
        pids = [p.get("pid") for p in etree.fromstring(package.read("docProps/custom.xml")).iter(f"{CUSTOM}property")]
    assert len(pids) == len(set(pids))


def test_a_file_that_is_not_its_format_is_refused(tmp_path):
    path = tmp_path / "not.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(zipfile.BadZipFile):
        pm.mark_document(str(path), "docx", COMPOSITE)
    with pytest.raises(ValueError):
        pm.mark_document(str(path), "pptx", COMPOSITE)


# --------------------------------------------------------------------------- PNG and JPEG


def test_a_png_carries_the_mark_and_its_pixels_are_untouched():
    data = _image("PNG")
    marked = pm.marked_image_bytes(data, "image/png", TRAINED)

    _assert_mark(_png_packet(marked), TRAINED)
    assert _png_chunks(marked)[0][0] == b"IHDR" and _png_chunks(marked)[1][0] == b"iTXt"
    assert _pixels(marked) == _pixels(data)


def test_a_jpeg_carries_the_mark_and_its_pixels_are_untouched():
    data = _image("JPEG")
    marked = pm.marked_image_bytes(data, "image/jpeg", TRAINED)

    _assert_mark(_jpeg_packet(marked), TRAINED)
    assert _jpeg_segments(marked)[0][0] == 0xE0  # JFIF's APP0 stays first
    assert _pixels(marked) == _pixels(data)


def test_an_image_with_a_c2pa_manifest_is_kept_byte_for_byte():
    png = _png_with(b"caBX", b"\x00\x00\x00\x10jumbc2pa-manifest")
    jpeg = _jpeg_with(0xEB, b"JP\x00\x01\x00\x00\x00\x01\x00\x00\x00\x10jumbc2pa")
    assert pm.marked_image_bytes(png, "image/png", TRAINED) is png
    assert pm.marked_image_bytes(jpeg, "image/jpeg", TRAINED) is jpeg


def test_an_image_whose_xmp_states_its_source_keeps_it():
    theirs = pm.xmp_packet(TRAINED).replace(TRAINED.encode(), b"http://example.test/their-own-term")
    png = _png_with(b"iTXt", b"XML:com.adobe.xmp\x00\x00\x00\x00\x00" + theirs)
    assert pm.marked_image_bytes(png, "image/png", TRAINED) is png


def test_an_image_cut_short_is_refused_and_another_type_passes_unmarked():
    with pytest.raises(ValueError):
        pm.marked_image_bytes(_image("PNG")[:20], "image/png", TRAINED)
    with pytest.raises(ValueError):
        pm.marked_image_bytes(b"\xff\xd8\xff\xe0\x00", "image/jpeg", TRAINED)
    gif = _image("GIF")
    assert pm.marked_image_bytes(gif, "image/gif", TRAINED) is gif


# --------------------------------------------------------------------------- FR-16


@pytest.mark.parametrize("fmt, make", [("pdf", _pdf), ("docx", _docx), ("xlsx", _xlsx)])
def test_two_marks_of_one_document_never_share_a_hash(tmp_path, fmt, make):
    first = make(tmp_path / f"first.{fmt}")
    second = tmp_path / f"second.{fmt}"
    shutil.copyfile(first, second)
    assert _sha(first) == _sha(second)  # identical input

    assert pm.mark_document(str(first), fmt, COMPOSITE) != pm.mark_document(str(second), fmt, COMPOSITE)


@pytest.mark.parametrize("fmt, mime", [("PNG", "image/png"), ("JPEG", "image/jpeg")])
def test_two_marks_of_one_image_never_share_a_hash(fmt, mime):
    data = _image(fmt)
    first, second = (pm.marked_image_bytes(data, mime, TRAINED) for _ in range(2))
    assert pm.sha256_hex(first) != pm.sha256_hex(second)
