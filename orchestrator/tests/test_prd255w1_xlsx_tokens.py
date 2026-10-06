"""PRD-255 Wave 1, US-005: a generated spreadsheet reads the kit's tokens, like the kit's documents.

Before, ``xlsx_render`` filled every header row with the kit's primary ("a lot of
orange in there"), printed every cell at Excel's default size, put no currency on
an amount, and drew the logo 48 px tall whatever the kit said. A branded sheet
now takes, through ``services.brand_rules.brand_assets`` and
``blocks.design_tokens``:

* the header row on ``surface_2`` with ``heading`` text, or on the ``accent`` only
  when the kit's ``accent_use`` is ``bold``;
* every cell in the kit's body font at ``type_scale.body``;
* an amount column (its header's last word) in the kit's currency with two
  decimals; a kit without a currency keeps the plain number;
* the logo at ``logo_rules.letterhead_mm``, the table under it and its clear space;
* the company line's date in the kit's ``date_style``.

Each sheet is the real one ``write_xlsx`` writes, read back with openpyxl.
"""
from __future__ import annotations

import base64
import re
import struct
import zipfile
import zlib
from types import SimpleNamespace as NS
from uuid import UUID

import openpyxl
import pytest

from core.brand_palette import derive_palette
from modules.documents.blocks.design_tokens import palette
from modules.documents.brand_kit import get_brand_kit
from modules.documents.brand_system import DATE_STYLE_MONTH_FIRST
from modules.documents.xlsx_letterhead import DEFAULT_ROW_PX, LETTERHEAD_ROWS, PX_PER_MM, logo_rows, sheet_design
from modules.documents.xlsx_render import write_xlsx
from services.brand_rules import brand_assets, forget_cached_kits

WS = UUID("6d0b5c1e-8f1a-4c2b-9d3e-0a1b2c3d4e62")
TITLE = "Green coffee stock"
DATA = {"columns": ["Coffee", "Quantity", "Revenue"],
        "rows": [["Kirinyaga", 41, 4500.0], ["Guji Shakiso", 55, 3000.5], ["Huila", 12, 960.0]]}
HARBOURLINE = {"name": "Harbourline Coffee Roasters", "primary_color": "#e96235", "secondary_color": "#1d3658",
               "text_color": "#2b2118", "font_family": "Georgia, serif"}
EMU_PER_MM = 36000
EMU_PER_PX = 9525


def _png(side: int = 8) -> bytes:
    """A small opaque square PNG: the logo."""
    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)

    rows = b"".join(b"\x00" + b"\xe9\x62\x35" * side for _ in range(side))
    header = struct.pack(">IIBBBBB", side, side, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")


LOGO = "data:image/png;base64," + base64.b64encode(_png()).decode("ascii")


@pytest.fixture(autouse=True)
def fresh_kits():
    forget_cached_kits()
    yield
    forget_cached_kits()


def _assets(kit):
    """The brand dict ``generate_xlsx`` hands the writer, for a workspace storing ``kit``."""
    settings = {"brand_kit": kit}
    db = NS(get=lambda model, key: NS(settings=settings), query=lambda *a, **k: None)
    return brand_assets(db, WS)


def _sheet(tmp_path, brand):
    path = str(tmp_path / "sheet.xlsx")
    write_xlsx(path, TITLE, DATA, brand)
    return path, openpyxl.load_workbook(path).active


def _header_row(sheet):
    return next(row for row in sheet.iter_rows(min_col=1, max_col=1) if row[0].value == "Coffee")[0].row


def _argb(hex_colour):
    return "FF" + hex_colour.lstrip("#").upper()


def test_brand_assets_carry_the_kits_design_system():
    kit = {**HARBOURLINE, "currency": "GBP", "accent_use": "bold", "date_style": DATE_STYLE_MONTH_FIRST,
           "logo_rules": {"letterhead_mm": 24}, "type_scale": {"body": {"size_pt": 11}}}
    brand = _assets(kit)
    stored = get_brand_kit({"brand_kit": kit})
    assert brand["palette"] == derive_palette(stored)
    assert brand["accent_use"] == "bold" and brand["currency"] == "GBP"
    assert brand["date_style"] == DATE_STYLE_MONTH_FIRST
    assert brand["logo_rules"]["letterhead_mm"] == 24 and brand["type_scale"]["body"]["size_pt"] == 11


@pytest.mark.parametrize("accent_use", ["sparing", "bold"])
def test_the_header_row_is_the_kits_table_header_never_the_primary_under_sparing(tmp_path, accent_use):
    kit = {**HARBOURLINE, "accent_use": accent_use}
    roles = palette(get_brand_kit({"brand_kit": kit}))
    _path, sheet = _sheet(tmp_path, _assets(kit))
    header = sheet.cell(row=_header_row(sheet), column=1)
    expected = roles.accent if accent_use == "bold" else roles.surface_2
    assert header.fill.fgColor.rgb == _argb(expected)
    assert header.font.color.rgb == _argb(roles.header_text)
    if accent_use == "sparing":
        assert roles.header_text == roles.heading
        assert header.fill.fgColor.rgb not in {_argb(HARBOURLINE["primary_color"]), _argb(roles.accent)}


def test_the_zebra_rows_are_on_the_kits_surface(tmp_path):
    roles = palette(get_brand_kit({"brand_kit": HARBOURLINE}))
    _path, sheet = _sheet(tmp_path, _assets(HARBOURLINE))
    top = _header_row(sheet)
    assert sheet.cell(row=top + 2, column=1).fill.fgColor.rgb == _argb(roles.surface)
    assert sheet.cell(row=top + 1, column=1).fill.fill_type is None


def test_every_cell_is_in_the_kits_body_font_and_size(tmp_path):
    kit = {**HARBOURLINE, "type_scale": {"body": {"size_pt": 11, "line_pt": 16}}}
    _path, sheet = _sheet(tmp_path, _assets(kit))
    top = _header_row(sheet)
    cells = [sheet.cell(row=top + r, column=c) for r in range(0, 4) for c in range(1, 4)]
    assert {cell.font.name for cell in cells} == {"Georgia"}
    assert {cell.font.sz for cell in cells} == {11}


def test_the_default_kit_prints_the_default_body_size(tmp_path):
    body = sheet_design({"colours": {}}).type["body"].size_pt
    _path, sheet = _sheet(tmp_path, _assets(HARBOURLINE))
    top = _header_row(sheet)
    assert sheet.cell(row=top + 1, column=1).font.sz == body
    assert sheet.cell(row=top, column=1).font.sz == body


def test_an_amount_column_prints_in_the_kits_currency_with_two_decimals(tmp_path):
    _path, sheet = _sheet(tmp_path, _assets({**HARBOURLINE, "currency": "GBP"}))
    top = _header_row(sheet)
    revenue = [sheet.cell(row=top + r, column=3) for r in (1, 2, 3)]
    assert {cell.number_format for cell in revenue} == {'"£"#,##0.00'}
    assert [cell.value for cell in revenue] == [4500, 3000.5, 960]
    quantity = sheet.cell(row=top + 1, column=2)
    assert "£" not in quantity.number_format    # a quantity is not money


def test_a_kit_without_a_currency_keeps_the_plain_number(tmp_path):
    path, sheet = _sheet(tmp_path, _assets(HARBOURLINE))
    top = _header_row(sheet)
    assert sheet.cell(row=top + 1, column=3).number_format == "General"
    assert sheet.cell(row=top + 2, column=3).number_format == "#,##0.00"
    with zipfile.ZipFile(path) as book:
        styles = book.read("xl/styles.xml").decode()
    assert "£" not in styles and "$" not in styles


@pytest.mark.parametrize("letterhead_mm", [16, 30])
def test_the_logo_is_the_kits_letterhead_height_and_the_table_sits_under_it(tmp_path, letterhead_mm):
    kit = {**HARBOURLINE, "logo_rules": {"letterhead_mm": letterhead_mm}}
    brand = {**_assets(kit), "logo": LOGO}
    path, sheet = _sheet(tmp_path, brand)
    with zipfile.ZipFile(path) as book:
        drawing = book.read("xl/drawings/drawing1.xml").decode()
    height_emu = int(re.search(r'<a:ext cx="\d+" cy="(\d+)"', drawing).group(1))
    assert abs(height_emu - letterhead_mm * EMU_PER_MM) <= EMU_PER_PX
    rows = logo_rows(sheet_design(brand))
    assert sheet.cell(row=rows + 1, column=1).value == TITLE
    assert _header_row(sheet) == rows + LETTERHEAD_ROWS + 1
    design = sheet_design(brand)
    assert design.logo_mm == letterhead_mm
    assert rows * DEFAULT_ROW_PX >= (design.logo_mm + design.logo_clear_mm) * PX_PER_MM    # logo + clear space fit


def test_the_company_line_dates_in_the_kits_date_style(tmp_path):
    _path, sheet = _sheet(tmp_path, _assets({**HARBOURLINE, "date_style": DATE_STYLE_MONTH_FIRST}))
    company_line = sheet.cell(row=2, column=1).value
    assert company_line.startswith(HARBOURLINE["name"])
    assert re.search(r"[A-Z][a-z]+ \d{1,2}, \d{4}$", company_line)


def test_without_a_kit_the_sheet_is_unchanged(tmp_path):
    _path, sheet = _sheet(tmp_path, None)
    header = sheet.cell(row=1, column=1)
    assert header.value == "Coffee" and header.fill.fgColor.rgb == "FF1A1A2E"
    assert sheet.cell(row=2, column=3).number_format == "General"
