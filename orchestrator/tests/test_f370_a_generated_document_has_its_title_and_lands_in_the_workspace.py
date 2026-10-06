"""F370 (night 10c, document parts): the spreadsheet wears data.title, and every document lands in documents/.

* The plain XLSX's sheet, letterhead title and printed header took the request's
  title ("Wholesale price list XLSX"), never ``data.title``. Now the data's title
  wins, and the sheet's name is cleaned of what Excel refuses ("2026/27").
* A generated PDF, Word file or spreadsheet lived only behind the API and object
  storage: the owner's deliverables folder (locally ~/Development/deliverables)
  never held it. When it is registered as a Deliverable it is now copied into
  ``documents/<file>`` through the workspace worker, as the Socials and session
  copies are; the copy is not registered twice.
"""
from __future__ import annotations

import asyncio
from pathlib import Path
from types import SimpleNamespace
from uuid import UUID

import openpyxl
import pytest

import modules.documents.generation_service as gs
import services.deliverable_service as ds
from modules.documents import workspace_documents
from modules.documents.xlsx_render import sheet_name, sheet_title, write_xlsx
from services.brand_rules import brand_assets, forget_cached_kits

WS = UUID("00000000-0000-0000-0000-0000000370a1")
PRICES = {"title": "Harbourline wholesale prices 2026/27", "columns": ["blend", "price"],
          "rows": [["Harbour Blend", 22.0], ["Quay Decaf", 24.5]]}


def test_the_sheet_takes_the_datas_title_over_the_requests():
    assert sheet_title("Wholesale price list XLSX", PRICES) == "Harbourline wholesale prices 2026/27"
    assert sheet_title("Wholesale price list XLSX", {"title": "  "}) == "Wholesale price list XLSX"
    assert sheet_name("Harbourline wholesale prices 2026/27") == "Harbourline wholesale prices 20"
    assert sheet_name("[Q3]: costs?") == "Q3   costs" and sheet_name("'/'") == "Sheet1"


def _brand():
    """The brand dict ``generate_xlsx`` hands the writer, for a workspace with a kit (as the W1 xlsx tests build it)."""
    settings = {"brand_kit": {"name": "Harbourline Coffee Roasters", "primary_color": "#e96235"}}
    forget_cached_kits()
    db = SimpleNamespace(get=lambda model, key: SimpleNamespace(settings=settings), query=lambda *a, **k: None)
    return brand_assets(db, WS)


def test_the_written_sheet_is_named_and_titled_from_the_data(tmp_path):
    path = tmp_path / "prices.xlsx"
    write_xlsx(str(path), "Wholesale price list XLSX", PRICES, _brand())
    forget_cached_kits()

    sheet = openpyxl.load_workbook(path).active
    assert sheet.title == sheet_name(PRICES["title"])
    texts = [str(cell.value) for row in sheet.iter_rows() for cell in row if cell.value is not None]
    assert PRICES["title"] in texts and "Wholesale price list XLSX" not in texts


class _Worker:
    written: dict = {}

    def __init__(self, workspace_id):
        self.workspace_id = workspace_id

    async def write_binary(self, path, pieces):
        self.written[(self.workspace_id, path)] = b"".join([piece async for piece in pieces])
        return {"success": True, "path": path}


@pytest.fixture
def generated(tmp_path, monkeypatch) -> Path:
    _Worker.written = {}
    monkeypatch.setattr(gs, "GENERATED_DIR", str(tmp_path))
    monkeypatch.setattr(workspace_documents, "WorkspaceClient", _Worker)
    folder = tmp_path / str(WS) / "generated"
    folder.mkdir(parents=True)
    (folder / "20261006_invoice.pdf").write_bytes(b"%PDF-1.7 invoice")
    (tmp_path / "secret.pdf").write_bytes(b"not yours")
    return folder


def test_a_generated_document_is_copied_into_documents(generated):
    assert asyncio.run(workspace_documents.copy_document(WS, "20261006_invoice.pdf")) == "documents/20261006_invoice.pdf"
    assert _Worker.written == {(str(WS), "documents/20261006_invoice.pdf"): b"%PDF-1.7 invoice"}


@pytest.mark.parametrize("name", ["../secret.pdf", "missing.pdf", "a/b.pdf", ".."])
def test_nothing_outside_the_workspaces_own_documents_is_copied(generated, name):
    assert asyncio.run(workspace_documents.copy_document(WS, name)) is None
    assert _Worker.written == {}


def test_only_documents_are_copied_and_only_from_a_running_loop(generated):
    async def started(fmt: str) -> bool:
        return workspace_documents.copy_when_registered(WS, "20261006_invoice.pdf", fmt)

    assert asyncio.run(started("pdf")) is True
    assert asyncio.run(started("social_image")) is False
    assert workspace_documents.copy_when_registered(WS, "20261006_invoice.pdf", "pdf") is False  # no loop: skipped


def test_registering_a_document_starts_its_copy(monkeypatch):
    started = []
    monkeypatch.setattr(gs, "copy_when_registered", lambda ws, name, fmt: started.append((ws, name, fmt)))

    class _Deliverables:
        def __init__(self, db, workspace_id):
            pass

        def register(self, **kwargs):
            return {"success": True, "deliverable_id": "d-1"}

    monkeypatch.setattr(ds, "DeliverableService", _Deliverables)
    result = gs.GeneratedDocument(path="/x/20261006_invoice.pdf", format="pdf", filename="20261006_invoice.pdf", size=9)

    gs.DocumentGenerationService(object(), WS).register_as_deliverable(result, title="Invoice HL-W-1034")

    assert started == [(WS, "20261006_invoice.pdf", "pdf")]
