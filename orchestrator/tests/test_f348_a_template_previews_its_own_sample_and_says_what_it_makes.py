"""F348 (night 10b): a template previews its own sample, and says which files it makes.

* ``POST /api/documents/templates/<id>/preview`` answered 422 "template variables did
  not resolve" for every field when sent the template's own sample: the starters and
  the Studio store it as ``{"data": {...}}`` and the route took the body as the data.
* The "Meeting Notes" starter has no body (no blocks, no uploaded .docx) and its .docx
  failed with "DOCX generation requires either a block template or an uploaded .docx".
* An xlsx asked of a PDF template answered a bare 400 "Invalid document generation
  request"; nothing, the gallery included, said which formats the template makes.

Boundaries faked (the session, object storage, the template lookup); the renders are real.
"""
from __future__ import annotations

import asyncio
from datetime import datetime
from types import SimpleNamespace as NS
from uuid import uuid4

import pytest
from fastapi import HTTPException

WS = uuid4()
TEMPLATE_ID = uuid4()


class _Query:
    def filter(self, *args, **kwargs):
        return self

    def order_by(self, *args, **kwargs):
        return self

    def first(self):
        return None


class _Db:
    """A session with no rows."""

    def query(self, *args, **kwargs):
        return _Query()


def _template(**columns):
    base = {"id": TEMPLATE_ID, "name": "T", "format": "pdf", "blocks": None, "template_content": None,
            "template_file_path": None, "data_schema": None, "sample_data": {}}
    return NS(**{**base, **columns})


def _letter():
    from modules.documents.presets import LETTER

    return _template(name=LETTER["name"], format=LETTER["format"], blocks=LETTER["blocks"],
                     sample_data=LETTER["sample_data"])


def _meeting_notes():
    from modules.documents.seed_templates import STARTER_TEMPLATES

    row = next(t for t in STARTER_TEMPLATES if t["name"] == "Meeting Notes")
    return _template(name=row["name"], format=row["format"], data_schema=row["data_schema"],
                     sample_data=row["sample_data"])


def _basic_report():
    return _template(name="Basic Report", format="pdf", template_content="<h1>{{ title }}</h1>")


# --------------------------------------------------------------------------- #
# 1. A preview takes the stored sample's shape, and an empty body is the sample
# --------------------------------------------------------------------------- #


def _preview(monkeypatch, template, body):
    """Run the preview route; returns (answer, the data it rendered from)."""
    import api.document_generation as documents_api
    import modules.documents.generation_service as gs
    import modules.documents.template_service as ts

    rendered = []

    class _Templates:
        def __init__(self, db):
            pass

        def get_template(self, template_id, workspace_id):
            return template

    class _Generation:
        def __init__(self, db, workspace_id):
            pass

        async def generate(self, **kwargs):
            rendered.append(kwargs["data"])
            kwargs["data"]["title"] = kwargs["title"]  # generation writes into its data
            return NS(download_url="/api/documents/generated/p.pdf", filename="p.pdf")

    monkeypatch.setattr(ts, "DocumentTemplateService", _Templates)
    monkeypatch.setattr(gs, "DocumentGenerationService", _Generation)
    ctx = NS(workspace_id=WS, user=None)
    answer = asyncio.run(documents_api.preview_template(template_id=TEMPLATE_ID, data=body, ctx=ctx, db=_Db()))
    return answer, rendered


def test_the_templates_own_wrapped_sample_previews_as_its_data(monkeypatch):
    template = _letter()
    answer, rendered = _preview(monkeypatch, template, body=template.sample_data)

    assert answer == {"preview_url": "/api/documents/generated/p.pdf", "filename": "p.pdf"}
    assert rendered[0]["recipient_name"] == "Jordan Smith"
    assert "data" not in rendered[0]


@pytest.mark.parametrize("body", [None, {}, {"data": {}}])
def test_an_empty_body_previews_the_stored_sample_and_never_writes_into_it(monkeypatch, body):
    template = _letter()
    stored = {"data": dict(template.sample_data["data"])}

    _, rendered = _preview(monkeypatch, template, body=body)

    assert rendered[0]["subject"] == "Your proposal for the spring campaign"
    assert template.sample_data == stored


def test_a_flat_body_still_previews_as_it_is(monkeypatch):
    _, rendered = _preview(monkeypatch, _letter(), body={"recipient_name": "Ana", "subject": "S", "body": "B"})

    assert rendered[0]["recipient_name"] == "Ana"


def test_the_wrapped_sample_fills_every_data_field_of_its_template():
    """The root cause, at the resolver: unwrapped, each ``data.*`` chip finds its value."""
    from modules.documents.blocks import collect_variable_paths, validate_blocks
    from modules.documents.template_preview import preview_data
    from modules.documents.variables.resolver import build_context, resolve_paths

    template = _letter()
    paths = [p for p in collect_variable_paths(validate_blocks(template.blocks)) if p.startswith("data.")]
    context = build_context(None, None, {}, datetime(2026, 10, 5), preview_data(template.sample_data, None))

    resolved = resolve_paths(context, paths)
    assert {p: resolved.values.get(p) for p in ("data.recipient_name", "data.subject")} == {
        "data.recipient_name": "Jordan Smith", "data.subject": "Your proposal for the spring campaign"}


# --------------------------------------------------------------------------- #
# 2. Meeting Notes makes its .docx
# --------------------------------------------------------------------------- #


def _service(monkeypatch, tmp_path, template):
    import modules.documents.generation_service as gs

    monkeypatch.setattr(gs, "GENERATED_DIR", str(tmp_path))
    monkeypatch.setattr(gs, "is_storage_configured", lambda: False)
    service = gs.DocumentGenerationService(_Db(), WS)
    service.template_service = NS(get_template=lambda template_id, ws: template,
                                  get_template_by_name=lambda ws, name: template)
    return service


def _docx_text(path):
    import docx

    document = docx.Document(path)
    cells = [cell.text for table in document.tables for row in table.rows for cell in row.cells]
    return "\n".join([p.text for p in document.paragraphs] + cells)


def test_the_meeting_notes_starter_makes_its_docx_from_its_sample(monkeypatch, tmp_path):
    template = _meeting_notes()
    service = _service(monkeypatch, tmp_path, template)

    result = asyncio.run(service.generate(title="Sprint Planning Meeting", format="docx",
                                          data=dict(template.sample_data), workspace_id=WS, template_id=TEMPLATE_ID))

    assert result.format == "docx" and result.filename.endswith(".docx")
    text = _docx_text(result.path)
    assert "Sprint Planning Meeting" in text
    assert "Alice" in text and "Complete API integration" in text


def test_the_meeting_notes_starter_says_it_makes_pdf_and_docx():
    from modules.documents.template_summary import summarize_template

    assert summarize_template(_meeting_notes())["supported_formats"] == ["pdf", "docx"]


# --------------------------------------------------------------------------- #
# 3. A format the template does not make is refused by name
# --------------------------------------------------------------------------- #


def test_xlsx_from_a_pdf_template_is_refused_naming_the_formats_it_makes(monkeypatch, tmp_path):
    from modules.documents.template_formats import UnsupportedTemplateFormat

    service = _service(monkeypatch, tmp_path, _letter())

    with pytest.raises(UnsupportedTemplateFormat) as refused:
        asyncio.run(service.generate(title="Letter", format="xlsx", data={"recipient_name": "Ana"},
                                     workspace_id=WS, template_id=TEMPLATE_ID))

    message = str(refused.value)
    assert "'Branded Letter' makes pdf or docx files, not xlsx" in message
    assert "'columns' and 'rows'" in message
    assert refused.value.supported == ["pdf", "docx"]
    assert list(tmp_path.rglob("*.xlsx")) == []


def test_docx_from_a_legacy_html_template_is_refused_naming_pdf(monkeypatch, tmp_path):
    from modules.documents.template_formats import UnsupportedTemplateFormat

    service = _service(monkeypatch, tmp_path, _basic_report())

    with pytest.raises(UnsupportedTemplateFormat, match="'Basic Report' makes pdf files, not docx"):
        asyncio.run(service.generate(title="R", format="docx", data={"title": "R"}, workspace_id=WS,
                                     template_id=TEMPLATE_ID))


def test_the_generate_route_answers_the_refusal_not_a_bare_invalid_request(monkeypatch):
    import api.document_generation as documents_api
    import modules.documents.generation_service as gs
    from modules.documents.template_formats import UnsupportedTemplateFormat

    class _Generation:
        def __init__(self, db, workspace_id):
            pass

        async def generate(self, **kwargs):
            raise UnsupportedTemplateFormat("Branded Report", "xlsx", ["pdf", "docx"])

    monkeypatch.setattr(gs, "DocumentGenerationService", _Generation)
    body = documents_api.GenerateDocumentRequest(title="R", format="xlsx", data={}, template_id=str(TEMPLATE_ID))

    with pytest.raises(HTTPException) as answered:
        asyncio.run(documents_api.generate_document(body=body, ctx=NS(workspace_id=WS, user=None), db=_Db()))

    assert answered.value.status_code == 400
    assert "'Branded Report' makes pdf or docx files, not xlsx" in answered.value.detail
    assert answered.value.detail != "Invalid document generation request"


@pytest.mark.parametrize("template, formats", [
    (_template(blocks={"version": 1, "blocks": []}), ["pdf", "docx"]),
    (_basic_report(), ["pdf"]),
    (_template(format="docx", template_file_path="/templates/letter.docx"), ["docx"]),
    (_template(format="xlsx"), ["xlsx"]),
    (_template(format="social_image", blocks={"variables_schema": {}}), ["social_image"]),
])
def test_the_template_summary_lists_the_formats_it_makes(template, formats):
    from modules.documents.template_summary import summarize_template

    assert summarize_template(template)["supported_formats"] == formats


def test_a_template_is_never_refused_its_own_format_and_no_template_is_never_refused():
    from modules.documents.template_formats import refuse_unsupported_format

    for template in (_letter(), _basic_report(), _meeting_notes(), _template(format="xlsx")):
        refuse_unsupported_format(template, template.format)
    refuse_unsupported_format(None, "xlsx")


def test_the_agents_template_list_says_what_each_template_makes(monkeypatch):
    import modules.documents.template_service as ts
    from modules.tools.discovery import handlers_documents

    rows = [_basic_report(), _letter(), _meeting_notes()]

    class _Templates:
        def __init__(self, db):
            pass

        def list_templates(self, workspace_id, format=None, category=None):
            return [NS(**{**vars(row), "description": None, "category": "report"}) for row in rows]

    monkeypatch.setattr(ts, "DocumentTemplateService", _Templates)
    answer = asyncio.run(handlers_documents.list_templates(_Db(), WS, {}))

    assert [(t["name"], t["supported_formats"]) for t in answer["templates"]] == [
        ("Basic Report", ["pdf"]), ("Branded Letter", ["pdf", "docx"]), ("Meeting Notes", ["pdf", "docx"])]
