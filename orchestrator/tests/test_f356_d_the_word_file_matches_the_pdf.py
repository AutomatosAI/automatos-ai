"""F356 (5 Oct): a Branded template's Word file matches its PDF.

The .docx was python-docx's default template: Calibri headings in Word's theme
blue, grid tables in "Light Grid Accent 1", no letterhead, no footer, no page
numbers. It now has the PDF's design system: the title in the kit's primary and
section headings in its accent, the letterhead as the first page's header, a
footer with the company, the title and Word page-number fields on every page,
tables with a primary header row, zebra rows and aligned figures, the report's
KPI tiles, and the Agreement's sign-off kept together.
"""
from __future__ import annotations

import base64
import copy
import struct
import zlib
from datetime import datetime
from types import SimpleNamespace
from typing import Any, Dict, List

import pytest

from modules.documents.blocks import collect_variable_paths, render_document_docx, validate_blocks
from modules.documents.brand_kit import get_brand_kit
from modules.documents.presets import preset_for
from modules.documents.variables.resolver import build_context, resolve_paths

pytest.importorskip("docx")

from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402

ORANGE, NAVY = "C44A1A", "1D3658"
COMPANY = "Automatos AI"
NOW = datetime(2026, 10, 5, 9, 0, 0)
USER = SimpleNamespace(name="Gerard Kavanagh", email="gerard@automatos.app", username="gerard")


def _png(side: int = 8) -> bytes:
    """A small opaque PNG: the logo, inlined as the renderer receives it."""
    def chunk(kind: bytes, body: bytes) -> bytes:
        return struct.pack(">I", len(body)) + kind + body + struct.pack(">I", zlib.crc32(kind + body) & 0xFFFFFFFF)

    rows = b"".join(b"\x00" + b"\xc4\x4a\x1a" * side for _ in range(side))
    header = struct.pack(">IIBBBBB", side, side, 8, 2, 0, 0, 0)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(rows)) + chunk(b"IEND", b"")


KIT = {
    **get_brand_kit({"brand_kit": {
        "name": COMPANY, "primary_color": f"#{ORANGE}", "accent_color": f"#{NAVY}",
        "company": {"name": COMPANY, "address": "14 Wapping Quay, Bristol BS1 4RW", "email": "gerard@automatos.app"},
    }}),
    "logo_url": "data:image/png;base64," + base64.b64encode(_png()).decode("ascii"),
}


def _docx(category: str, data: Dict[str, Any] = None) -> Any:
    preset = preset_for(category)
    data = copy.deepcopy(preset["sample_data"]["data"]) if data is None else data
    doc = validate_blocks(preset["blocks"])
    resolved = resolve_paths(build_context(USER, None, KIT, NOW, extra_data=data), collect_variable_paths(doc))
    rendered = render_document_docx(doc, resolved.values, KIT, data=data)
    assert rendered.unresolved == []
    return rendered.document


def _fills(row: Any) -> List[str]:
    return [shd.get(qn("w:fill")) for cell in row.cells for shd in cell._tc.iter(qn("w:shd"))]


def _texts(document: Any) -> List[str]:
    return [p.text for p in document.paragraphs]


def test_the_styles_take_the_kits_colours_and_the_type_scale():
    styles = _docx("report").styles
    assert str(styles["Heading 1"].font.color.rgb) == ORANGE and styles["Heading 1"].font.size.pt == 22
    assert str(styles["Heading 2"].font.color.rgb) == NAVY and styles["Heading 2"].font.size.pt == 13
    assert styles["Normal"].font.size.pt == 10
    fonts = styles["Heading 1"].element.rPr.find(qn("w:rFonts"))
    assert fonts is None or fonts.get(qn("w:asciiTheme")) is None  # Word's theme font no longer wins


def test_the_letterhead_is_the_first_pages_header():
    document = _docx("invoice")
    section = document.sections[0]
    assert section.different_first_page_header_footer
    header = section.first_page_header
    (table,) = header.tables
    logo_cell, company_cell = table.rows[0].cells
    assert logo_cell._tc.findall(".//" + qn("w:drawing"))
    assert company_cell.paragraphs[0].text == COMPANY
    assert all(p.alignment == WD_ALIGN_PARAGRAPH.RIGHT for p in company_cell.paragraphs)
    assert COMPANY not in _texts(document)  # the letterhead is not repeated in the body


def test_every_page_has_a_footer_with_page_numbers():
    section = _docx("invoice").sections[0]
    for footer in (section.footer, section.first_page_footer):
        xml = footer._element.xml
        assert 'w:instr="PAGE"' in xml and 'w:instr="NUMPAGES"' in xml
        text = footer.paragraphs[0].text
        assert COMPANY in text and "Invoice INV-0042" in text


def test_a_data_table_has_a_primary_header_zebra_rows_and_right_aligned_figures():
    document = _docx("invoice")
    items = document.tables[0]
    assert items.style is None or "Grid" not in items.style.name
    header, first, second = items.rows[0], items.rows[1], items.rows[2]
    assert set(_fills(header)) == {ORANGE}
    assert header._tr.trPr.find(qn("w:tblHeader")) is not None
    assert not _fills(first) and _fills(second)  # zebra: every second body row is tinted
    assert first.cells[3].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert first.cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_the_invoice_totals_end_bold_over_a_rule():
    totals = _docx("invoice").tables[1]
    last = totals.rows[-1]
    assert [cell.text for cell in last.cells] == ["Total due", "6,273.00"]
    assert all(run.bold for cell in last.cells for p in cell.paragraphs for run in p.runs)
    assert last.cells[0]._tc.find(".//" + qn("w:top")) is not None


def test_the_report_prints_its_kpis_as_tiles():
    tiles = _docx("report").tables[0]
    assert len(tiles.rows) == 1
    assert [cell.paragraphs[1].text for cell in tiles.rows[0].cells] == ["€182k", "312", "61"]


def test_the_agreements_sign_off_is_kept_together():
    document = _docx("contract")
    kept = {p.text: p.paragraph_format.keep_with_next for p in document.paragraphs}
    assert kept["6. Governing law"] and kept["Signed"]
    assert kept["This agreement is governed by the laws of Ireland."]
    signatures = document.tables[-1]
    assert all(row._tr.trPr.find(qn("w:cantSplit")) is not None for row in signatures.rows)


def test_an_optional_part_not_sent_is_left_out_of_the_word_file_too():
    data = {k: v for k, v in preset_for("proposal")["sample_data"]["data"].items() if k not in ("pricing_total", "subtitle")}
    document = _docx("proposal", data)
    assert not any(cell.text == "Total" for table in document.tables for row in table.rows for cell in row.cells)
    assert "" not in _texts(document)[:4]  # no empty line where the subtitle would be
