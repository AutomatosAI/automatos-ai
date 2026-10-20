"""PRD-262 US-025: pandas leaves the image.

- create_xlsx (``exec_document.execute_document_tool``) turns a CSV or JSON file in the
  workspace into a spreadsheet with the standard library and xlsxwriter, with pandas
  unimportable. The file is read back with openpyxl: the sheet's name, the header,
  numbers typed per column, text, blanks and booleans. The refusals are unchanged, and
  every other tool still goes through the function it decorates.
- PandasAIService binds pandas only once pandasai imports. With neither installed it
  reports itself unavailable, and importing its module never needs pandas.
- POST /api/analytics/charts/generate keeps its 404 "PandasAI service is not
  available…" with pandas and pandasai unimportable: never a 500 or an ImportError.
- pytz, which services/heartbeat_service.py imports and only pandas brought, stays.

Deleting the PandasAI service and its routes instead is §9-Q16, still open.
"""
from __future__ import annotations

import asyncio
import importlib.util
import json
import os
import re
import subprocess
import sys
import types
import uuid
from pathlib import Path
from types import SimpleNamespace as NS
from unittest.mock import MagicMock

import pytest

# Dummy POSTGRES_* satisfies the config chain at import (the pattern of
# test_prd143_obs_routers_batch1.py); CI exports real ones and setdefault no-ops.
os.environ.setdefault("POSTGRES_USER", "test")
os.environ.setdefault("POSTGRES_PASSWORD", "test")
os.environ.setdefault("POSTGRES_HOST", "127.0.0.1")
os.environ.setdefault("POSTGRES_PORT", "59432")
os.environ.setdefault("POSTGRES_DB", "test")

ORCH = Path(__file__).resolve().parents[1]
SERVICE_FILE = ORCH / "modules" / "tools" / "services" / "pandas_ai_service.py"
SERVICE_MODULE = "modules.tools.services.pandas_ai_service"
NOT_AVAILABLE = "PandasAI service is not available. Ensure pandasai is installed and configured."


@pytest.fixture
def no_pandas(monkeypatch):
    """pandas and pandasai unimportable, whatever this environment installs."""
    for name in ("pandas", "pandasai", "pandasai_litellm"):
        monkeypatch.setitem(sys.modules, name, None)


def _create_xlsx(workspace: Path, parameters: dict, tool_name: str = "create_xlsx") -> dict:
    from modules.tools.execution import exec_document

    executor = NS(workspace_dir=str(workspace))
    return asyncio.run(exec_document.execute_document_tool(executor, tool_name, parameters, 25))


def _rows(path: Path, sheet: str) -> list:
    import openpyxl

    book = openpyxl.load_workbook(path)
    assert book.sheetnames == [sheet]
    return list(book[sheet].iter_rows(values_only=True))


# --------------------------------------------------------------------------- create_xlsx


CSV = (
    "name,qty,price,note,code\r\n"
    "Kirinyaga,41,4500.5,washed,007\r\n"
    "\r\n"
    "Guji,,3000,,A1\r\n"
    'Huila,12,960,"natural, sweet",\r\n'
)


def test_a_csv_becomes_a_sheet_with_its_columns_typed_and_blanks_blank(tmp_path, no_pandas):
    (tmp_path / "lots.csv").write_text(CSV, encoding="utf-8-sig")  # with a byte-order mark
    (tmp_path / "out").mkdir()
    result = _create_xlsx(tmp_path, {"source_file": "lots.csv", "output_file": "out/lots.xlsx",
                                     "sheet_name": "Lot list"})
    output = str((tmp_path / "out" / "lots.xlsx").resolve())
    assert result["success"] is True and result["action"] == "create_xlsx", result
    assert result["output_file"] == output
    assert result["result"] == f"Excel file created successfully: {output}"
    assert _rows(Path(output), "Lot list") == [
        ("name", "qty", "price", "note", "code"),
        ("Kirinyaga", 41, 4500.5, "washed", "007"),  # code mixes text in: it stays text
        ("Guji", None, 3000, None, "A1"),  # the empty line is skipped, as pandas skipped it
        ("Huila", 12, 960, "natural, sweet", None),
    ]


def test_json_records_become_a_sheet_in_first_seen_column_order(tmp_path, no_pandas):
    records = [{"lot": "A", "kg": 12.5, "graded": True, "tags": ["washed"]},
               {"lot": "B", "kg": None, "graded": False, "origin": "Huila"}]
    (tmp_path / "lots.json").write_text(json.dumps(records), encoding="utf-8")
    result = _create_xlsx(tmp_path, {"source_file": "lots.json", "output_file": "lots.xlsx"})
    assert result["success"] is True, result
    assert _rows(tmp_path / "lots.xlsx", "Sheet1") == [
        ("lot", "kg", "graded", "tags", "origin"),
        ("A", 12.5, True, '["washed"]', None),
        ("B", None, False, None, "Huila"),
    ]


@pytest.mark.parametrize("columns", [
    {"lot": ["A", "B"], "kg": [1, 2.5]},
    {"lot": {"0": "A", "1": "B"}, "kg": {"0": 1, "1": 2.5}},
], ids=["lists", "keyed-by-row"])
def test_a_json_object_of_columns_becomes_a_sheet(tmp_path, no_pandas, columns):
    (tmp_path / "cols.json").write_text(json.dumps(columns), encoding="utf-8")
    result = _create_xlsx(tmp_path, {"source_file": "cols.json", "output_file": "cols.xlsx", "sheet_name": "Kg"})
    assert result["success"] is True, result
    assert _rows(tmp_path / "cols.xlsx", "Kg") == [("lot", "kg"), ("A", 1), ("B", 2.5)]


@pytest.mark.parametrize("name, text, error", [
    ("bad.json", "42", "Document creation failed: JSON must be a list of objects (one per row) "
                       "or an object of columns"),
    ("broken.json", "[{", "Document creation failed: "),
    ("ragged.csv", "a,b\n1,2,3\n", "Document creation failed: Data row 1 has 3 fields; the header has 2"),
    ("empty.csv", "", "Document creation failed: No columns to parse from file"),
    ("notes.txt", "a,b\n", "Source file must be CSV or JSON for XLSX creation"),
])
def test_a_source_it_cannot_read_fails_clearly_and_writes_nothing(tmp_path, no_pandas, name, text, error):
    (tmp_path / name).write_text(text, encoding="utf-8")
    result = _create_xlsx(tmp_path, {"source_file": name, "output_file": "out.xlsx"})
    assert result["success"] is False and result["tool"] == "create_xlsx", result
    assert result["error"].startswith(error), result["error"]
    assert not (tmp_path / "out.xlsx").exists()


def test_a_missing_source_file_fails_clearly(tmp_path, no_pandas):
    result = _create_xlsx(tmp_path, {"source_file": "absent.csv", "output_file": "out.xlsx"})
    assert result["success"] is False and result["error"].startswith("Document creation failed: "), result


def test_an_output_folder_that_does_not_exist_fails_clearly(tmp_path, no_pandas):
    """As pandas did, the tool makes no folders; xlsxwriter's own error is caught like an OSError."""
    (tmp_path / "lots.csv").write_text("a\n1\n", encoding="utf-8")
    result = _create_xlsx(tmp_path, {"source_file": "lots.csv", "output_file": "missing/out.xlsx"})
    assert result["success"] is False and result["error"].startswith("Document creation failed: "), result
    assert not (tmp_path / "missing").exists()


@pytest.mark.parametrize("parameters, error", [
    ({"output_file": "out.xlsx"}, "Missing required parameters: source_file and output_file"),
    ({"source_file": "lots.csv", "output_file": ""}, "Missing required parameters: source_file and output_file"),
    ({"source_file": "../lots.csv", "output_file": "out.xlsx"}, "File paths must be within the workspace directory"),
    ({"source_file": "lots.csv", "output_file": "/tmp/prd262-out.xlsx"},
     "File paths must be within the workspace directory"),
])
def test_the_document_tools_refusals_still_answer_create_xlsx(tmp_path, no_pandas, parameters, error):
    workspace = tmp_path / "workspace"
    workspace.mkdir()
    (tmp_path / "lots.csv").write_text("a\n1\n", encoding="utf-8")
    (workspace / "lots.csv").write_text("a\n1\n", encoding="utf-8")
    assert _create_xlsx(workspace, parameters) == {"success": False, "error": error, "tool": "create_xlsx"}


def test_another_document_tool_still_runs_the_function_it_decorates(tmp_path, monkeypatch):
    calls = []

    def run(cmd, **kwargs):
        calls.append(cmd)
        return NS(returncode=0, stderr="")

    monkeypatch.setattr(subprocess, "run", run)
    (tmp_path / "notes.md").write_text("# Notes\n", encoding="utf-8")
    result = _create_xlsx(tmp_path, {"source_file": "notes.md", "output_file": "notes.pdf"}, tool_name="create_pdf")
    assert result["success"] is True and result["action"] == "create_pdf", result
    assert calls and calls[0][0] == "pandoc"


# --------------------------------------------------------------------------- the PandasAI service


def _fresh_service_module() -> types.ModuleType:
    """pandas_ai_service.py loaded from its file (sibling tests stub the package path)."""
    spec = importlib.util.spec_from_file_location("prd262_pandas_ai_service", SERVICE_FILE)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_the_service_imports_without_pandas_and_reports_itself_unavailable(no_pandas):
    module = _fresh_service_module()
    assert module.pd is None
    assert module.get_pandasai_service() is None
    assert module.pd is None


def test_an_enabled_service_binds_pandas_for_its_insights(monkeypatch, tmp_path):
    from config import config

    pandas = types.ModuleType("pandas")
    pandasai = types.ModuleType("pandasai")
    pandasai.SmartDataframe = object
    pandasai_config = types.ModuleType("pandasai.config")
    pandasai_config.Config = lambda **kwargs: kwargs
    litellm_pkg = types.ModuleType("pandasai_litellm")
    litellm = types.ModuleType("pandasai_litellm.litellm")
    litellm.LiteLLM = lambda **kwargs: kwargs
    fakes = {"pandas": pandas, "pandasai": pandasai, "pandasai.config": pandasai_config,
             "pandasai_litellm": litellm_pkg, "pandasai_litellm.litellm": litellm}
    for name, fake in fakes.items():
        monkeypatch.setitem(sys.modules, name, fake)
    monkeypatch.setattr(config, "PANDASAI_API_KEY", "fake-key")
    monkeypatch.setattr(config, "PANDASAI_OUTPUT_DIR", str(tmp_path / "charts"))
    module = _fresh_service_module()
    assert module.pd is None, "loading the module binds nothing"
    service = module.PandasAIService()
    assert service.enabled is True
    assert module.pd is pandas, "generate_insight reads pandas as the module's pd"


# --------------------------------------------------------------------------- the charts route


def test_the_charts_route_keeps_its_404_without_pandas(monkeypatch, no_pandas):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from core.auth.dependencies import RequestContext, UserContext
    from core.auth.hybrid import get_request_context_hybrid
    from core.database.database import get_db

    import api.analytics_charts as charts

    monkeypatch.setitem(sys.modules, SERVICE_MODULE, _fresh_service_module())
    super_admin = UserContext(id="u-prd262", role="admin", system_role="super_admin")
    app = FastAPI()
    app.include_router(charts.router)

    def _db():
        yield MagicMock()

    app.dependency_overrides[get_request_context_hybrid] = lambda: RequestContext(
        workspace_id=uuid.uuid4(), user=super_admin, auth_type="clerk")
    app.dependency_overrides[get_db] = _db
    response = TestClient(app, raise_server_exceptions=False).post(
        "/api/analytics/charts/generate", json={"query": "cost by model"})
    assert response.status_code == 404, response.text
    assert response.json()["detail"] == NOT_AVAILABLE


# --------------------------------------------------------------------------- requirements


def _names(path: Path) -> set:
    return {re.sub(r"[-_.]+", "-", match.group(1)).lower()
            for line in path.read_text(encoding="utf-8").splitlines()
            if (match := re.match(r"^([A-Za-z0-9][A-Za-z0-9._-]*)", line.split("#", 1)[0].strip()))}


def test_pytz_stays_for_the_heartbeat_once_pandas_goes():
    """apscheduler 3.11 needs no pytz; heartbeat_service's active hours import it, and only pandas brought it."""
    heartbeat = (ORCH / "services" / "heartbeat_service.py").read_text(encoding="utf-8")
    assert re.search(r"^\s*import pytz\b", heartbeat, re.M), "drop pytz from requirements.in with its last import"
    for name in ("requirements.in", "requirements.lock"):
        names = _names(ORCH / name)
        assert "pytz" in names and "pandas" not in names, name
