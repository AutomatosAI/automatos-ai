"""F353 (issue #947): a Word document, a sheet, a CSV or markdown as one printable page.

These formats have no page of their own to draw, so their opening is laid out as
HTML (then printed by WeasyPrint and drawn like any PDF, see ``render``):

* a Word document: its headings, paragraphs and tables, in order (text only);
* a spreadsheet (.xlsx) or CSV/TSV: the header row and the first rows, as a table;
* markdown: the text as the report view shows it (F374); plain text as written.

Only the opening is read (``MAX_*`` below): the picture shows page 1, never more.
Every value is HTML-escaped; markdown goes through the sanitising renderer.
"""
from __future__ import annotations

import csv
import html
import io
from itertools import islice
from typing import Any, Callable, Dict, List

MAX_TABLE_ROWS = 25
MAX_TABLE_COLS = 8
MAX_CELL_CHARS = 60
MAX_TEXT_BYTES = 64 * 1024
MAX_DOCX_BLOCKS = 60

HEADING_TAGS = {"title": "h1", "heading 1": "h1", "heading 2": "h2"}

PAGE_CSS = """
@page { size: A4; margin: 14mm; }
body { font-family: sans-serif; font-size: 11pt; color: #111827; line-height: 1.4; }
h1 { font-size: 20pt; margin: 0 0 8pt; }
h2 { font-size: 15pt; margin: 10pt 0 6pt; }
h3 { font-size: 12pt; margin: 8pt 0 4pt; }
p { margin: 0 0 6pt; }
table { border-collapse: collapse; width: 100%; font-size: 9pt; margin: 4pt 0 8pt; }
th, td { border: 1px solid #cbd5e1; padding: 3pt 5pt; text-align: left; vertical-align: top; }
th { background: #eef2f7; font-weight: 600; }
pre { white-space: pre-wrap; font-size: 9pt; }
"""

# F374: markdown is drawn as the report view (frontend .md-view) shows it: headings at the
# view's sizes against its 14px body, inline code a small chip inside its line, lists indented.
MARKDOWN_CLASS = "md"
MARKDOWN_CSS = """
.md > * + * { margin-top: 0.9em; }
.md > :first-child { margin-top: 0; }
.md p, .md ul, .md ol, .md pre, .md table { margin-bottom: 0; }
.md h1, .md h2, .md h3, .md h4, .md h5, .md h6 {
  font-weight: 600; line-height: 1.25; margin: 1.7em 0 0;
}
.md h1 { font-size: 1.714em; padding-bottom: 0.3em; border-bottom: 1px solid #e5e7eb; }
.md h2 { font-size: 1.429em; padding-bottom: 0.25em; border-bottom: 1px solid #eef0f3; }
.md h3 { font-size: 1.214em; }
.md h4 { font-size: 1.071em; }
.md h5, .md h6 { font-size: 1em; color: #6b7280; }
.md ul { list-style: disc; padding-left: 1.35em; }
.md ol { list-style: decimal; padding-left: 1.45em; }
.md li { margin: 0; }
.md li + li, .md li > ul, .md li > ol { margin-top: 0.35em; }
.md ul.task-list { list-style: none; padding-left: 0.2em; }
.md code {
  font-family: monospace; font-size: 0.875em; font-weight: normal;
  padding: 0.1em 0.35em; border-radius: 3px; background: #f3f4f6; border: 1px solid #e5e7eb;
}
.md pre code { background: none; border: 0; padding: 0; }
.md blockquote { border-left: 3px solid #cbd5e1; padding-left: 1em; margin-left: 0; color: #6b7280; }
.md hr { border: 0; border-top: 1px solid #e5e7eb; }
"""

Rows = List[List[str]]


def page_html(body: str) -> str:
    """A whole A4 page around ``body``."""
    return (
        "<!doctype html><html><head><meta charset='utf-8'>"
        f"<style>{PAGE_CSS}{MARKDOWN_CSS}</style></head><body>{body}</body></html>"
    )


def _cell(value: Any) -> str:
    return "" if value is None else str(value)[:MAX_CELL_CHARS]


def table_html(rows: Rows) -> str:
    """The first row as the header, the rest as the body; every value escaped."""
    if not rows:
        return "<p>(empty)</p>"
    head, *body = rows
    header = "".join(f"<th>{html.escape(c)}</th>" for c in head)
    lines = "".join(
        "<tr>" + "".join(f"<td>{html.escape(c)}</td>" for c in row) + "</tr>" for row in body
    )
    return f"<table><thead><tr>{header}</tr></thead><tbody>{lines}</tbody></table>"


def _text(data: bytes) -> str:
    return data[:MAX_TEXT_BYTES].decode("utf-8-sig", errors="replace")


def xlsx_rows(data: bytes) -> Rows:
    """The active sheet's opening rows (values, not formulas)."""
    import openpyxl

    book = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        sheet = book.active
        if sheet is None:
            return []
        rows = sheet.iter_rows(max_row=MAX_TABLE_ROWS, max_col=MAX_TABLE_COLS, values_only=True)
        return [[_cell(v) for v in row] for row in rows]
    finally:
        book.close()


def delimited_rows(data: bytes, delimiter: str) -> Rows:
    """A CSV's (or TSV's) opening rows."""
    reader = csv.reader(io.StringIO(_text(data)), delimiter=delimiter)
    return [[_cell(v) for v in row[:MAX_TABLE_COLS]] for row in islice(reader, MAX_TABLE_ROWS)]


def _paragraph_html(paragraph: Any) -> str:
    text = paragraph.text.strip()
    if not text:
        return ""
    style = (getattr(paragraph.style, "name", "") or "").lower()
    tag = HEADING_TAGS.get(style) or ("h3" if style.startswith("heading") else "p")
    return f"<{tag}>{html.escape(text)}</{tag}>"


def _docx_table_rows(table: Any) -> Rows:
    return [[_cell(c.text) for c in row.cells[:MAX_TABLE_COLS]] for row in islice(table.rows, MAX_TABLE_ROWS)]


def docx_body(data: bytes) -> str:
    """A Word document's opening blocks, paragraphs and tables in their order."""
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = Document(io.BytesIO(data))
    parts: List[str] = []
    for child in islice(document.element.body.iterchildren(), MAX_DOCX_BLOCKS):
        tag = str(child.tag).rsplit("}", 1)[-1]
        if tag == "p":
            parts.append(_paragraph_html(Paragraph(child, document)))
        elif tag == "tbl":
            parts.append(table_html(_docx_table_rows(Table(child, document))))
    return "".join(parts) or "<p>(empty)</p>"


def markdown_body(data: bytes) -> str:
    """Markdown as the report view shows it, sanitised (F374: the document renderer's
    reading of an agent's markdown, ``view_html``, not the blog's)."""
    from modules.documents.blocks.markdown_body import view_html

    return f'<div class="{MARKDOWN_CLASS}">{view_html(_text(data))}</div>'


def text_body(data: bytes) -> str:
    """Plain text, as written."""
    return f"<pre>{html.escape(_text(data))}</pre>"


BODY_BUILDERS: Dict[str, Callable[[bytes], str]] = {
    ".docx": docx_body,
    ".xlsx": lambda data: table_html(xlsx_rows(data)),
    ".csv": lambda data: table_html(delimited_rows(data, ",")),
    ".tsv": lambda data: table_html(delimited_rows(data, "\t")),
    ".md": markdown_body,
    ".markdown": markdown_body,
    ".txt": text_body,
}


def document_html(data: bytes, ext: str) -> str:
    """The page for a file of type ``ext`` (one of ``BODY_BUILDERS``)."""
    return page_html(BODY_BUILDERS[ext](data))
