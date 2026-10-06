"""The brand board's parts in a Word file (PRD-255 US-009).

The same parts as the PDF (``brand_board_html``), read from the kit by
``brand_board``: the logo with the name and tagline, the logo variants on their
grounds (an unset variant as its FR-9 fallback, saying so), the colour roles as
shaded swatches with their hex codes, the type scale with samples, the spacing,
the voice, and the applications (the invoice and letter miniatures as pictures,
the social card as a shaded cell in the kit's colours).

Images are read through the renderer's own SSRF-guarded loader, handed in as
``load_image`` (``docx_renderer._safe_image_bytes``).
"""
from __future__ import annotations

import logging
from typing import Any, Callable, Dict, Mapping, Optional

from . import brand_board as bb
from . import design_tokens as tokens
from .brand_board_style import BOARD_TOKENS
from .docx_style import full_width, rgb, text_width
from .docx_tables import shade

logger = logging.getLogger(__name__)

ImageLoader = Callable[[str], Optional[Any]]
# A miniature's width in the Word file (mm): three fit across the text column.
DOCX_APP_MM = 52


class _Ctx:
    """What every part writes with: the document, the kit's design, the font and the image loader."""

    def __init__(self, doc: Any, kit: Mapping[str, Any], font: Optional[str], load_image: ImageLoader) -> None:
        self.doc, self.kit, self.font, self.load_image = doc, kit, font, load_image
        self.design = tokens.design(kit)


def _run(paragraph: Any, ctx: _Ctx, text: str, step: str, colour: str, bold: Optional[bool] = None) -> Any:
    from docx.shared import Pt

    run = paragraph.add_run(text)
    found = ctx.design.type[step]
    run.font.size, run.font.color.rgb = Pt(found.size_pt), rgb(colour)
    run.bold = found.bold if bold is None else bold
    if ctx.font:
        run.font.name = ctx.font
    return run


def _line(container: Any, ctx: _Ctx, text: str, step: str = "caption", colour: Optional[str] = None,
          bold: Optional[bool] = None) -> Any:
    """A paragraph of ``text`` in ``step`` (in a cell's first, empty paragraph when there is one)."""
    from docx.shared import Pt

    paragraphs = getattr(container, "paragraphs", [])
    reuse = container is not ctx.doc and len(paragraphs) == 1 and not paragraphs[0].text
    paragraph = paragraphs[0] if reuse else container.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(ctx.design.space(1))
    _run(paragraph, ctx, text, step, colour or ctx.design.palette.muted, bold)
    return paragraph


def _label(ctx: _Ctx, text: str) -> None:
    _line(ctx.doc, ctx, text.upper(), "caption", ctx.design.palette.heading, True)


def _picture(cell: Any, ctx: _Ctx, src: str, **size: Any) -> bool:
    """``src`` as a picture in ``cell``; False when it cannot be loaded or embedded."""
    stream = ctx.load_image(src) if src else None
    if stream is None:
        return False
    try:
        cell.paragraphs[0].add_run().add_picture(stream, **size)
    except Exception:  # noqa: BLE001 — an unreadable image format: the part goes without it
        logger.warning("[BrandBoard] an image could not be embedded in the Word file; left out", exc_info=True)
        return False
    return True


def _table(ctx: _Ctx, rows: int, cols: int) -> Any:
    table = ctx.doc.add_table(rows=rows, cols=cols)
    full_width(table, text_width(ctx.design))
    return table


def logo_part(ctx: _Ctx) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm

    roles, name = ctx.design.palette, bb.brand_name(ctx.kit)
    table = _table(ctx, 1, 2)
    mark, words = table.cell(0, 0), table.cell(0, 1)
    if not _picture(mark, ctx, bb.logo(ctx.kit), height=Mm(BOARD_TOKENS["board_logo_mm"])):
        _line(mark, ctx, name, "display", roles.heading)
        _line(mark, ctx, bb.NO_LOGO_NOTE)
    lines = ((bb.BOARD_TITLE.upper(), "caption", roles.heading), (name, "h2", roles.heading),
             (bb.tagline(ctx.kit), "body", roles.muted))
    for text, step, colour in lines:
        if text:
            _line(words, ctx, text, step, colour).alignment = WD_ALIGN_PARAGRAPH.RIGHT


def variants_part(ctx: _Ctx) -> None:
    from docx.shared import Mm

    _label(ctx, "Logo")
    found = bb.variants(ctx.kit)
    table = _table(ctx, 2, len(found))
    for column, variant in enumerate(found):
        ground = table.cell(0, column)
        shade(ground, variant.ground)
        target = ground
        if variant.chip:  # FR-9: the logo on a light chip on the dark ground
            target = ground.add_table(rows=1, cols=1).cell(0, 0)
            shade(target, ctx.design.palette.paper)
        _picture(target, ctx, variant.src, height=Mm(BOARD_TOKENS["board_variant_logo_mm"]))
        caption = table.cell(1, column)
        _line(caption, ctx, variant.label, "small", ctx.design.palette.heading, True)
        if variant.note:
            _line(caption, ctx, variant.note)


def colours_part(ctx: _Ctx) -> None:
    _label(ctx, "Colour")
    found = bb.swatches(ctx.kit)
    table = _table(ctx, 2, len(found))
    for column, swatch in enumerate(found):
        shade(table.cell(0, column), swatch.hex)
        _line(table.cell(0, column), ctx, " ")
        cell = table.cell(1, column)
        _line(cell, ctx, swatch.label, "caption", ctx.design.palette.heading, True)
        _line(cell, ctx, swatch.hex, "caption", ctx.design.palette.ink)
    _line(ctx.doc, ctx, bb.accent_rule(ctx.kit))


def type_part(ctx: _Ctx) -> None:
    from docx.shared import Pt

    _label(ctx, "Type")
    for line in bb.font_lines(ctx.kit):  # F360: the fonts, and a substitute the PDF prints
        _line(ctx.doc, ctx, line.text)
    samples = bb.type_samples(ctx.kit)
    table = _table(ctx, len(samples), 2)
    for row, sample in enumerate(samples):
        _line(table.cell(row, 0), ctx, sample.label)
        paragraph = _line(table.cell(row, 1), ctx, bb.TYPE_SAMPLE, sample.step, ctx.design.palette.heading)
        paragraph.paragraph_format.line_spacing = Pt(sample.line_pt)


def spacing_part(ctx: _Ctx) -> None:
    grid = bb.spacing(ctx.kit)
    _label(ctx, "Spacing")
    gaps = ", ".join(f"{gap:g}" for gap in grid.gaps_pt)
    _line(ctx.doc, ctx, f"A {grid.unit_pt:g} pt grid: gaps of {gaps} pt; page margins {grid.margin_mm:g} mm.")
    _line(ctx.doc, ctx, f"Logo clear space: {grid.clear_space:.2g} of its height "
                        f"({grid.clear_mm:.3g} mm round the {grid.logo_mm:g} mm letterhead logo).")
    _label(ctx, bb.LOCALE_LABEL)  # F368: the kit's currency (or that it has none) and its date style
    for line in bb.locale_lines(ctx.kit):
        _line(ctx.doc, ctx, line.text)


def voice_part(ctx: _Ctx) -> None:
    roles = ctx.design.palette
    tones, sign_off = bb.voice(ctx.kit)
    _label(ctx, "Voice")
    for tone in tones:
        paragraph = _line(ctx.doc, ctx, tone["word"], "small", roles.heading, True)
        if tone["meaning"]:
            _run(paragraph, ctx, f"  {tone['meaning']}", "small", roles.ink, False)
    if not tones:
        _line(ctx.doc, ctx, bb.NO_TONE_NOTE)
    if sign_off:
        _line(ctx.doc, ctx, f"Signs off as {sign_off}")


def _social_cell(cell: Any, ctx: _Ctx) -> None:
    paper, ink, _accent = bb.social_colours(ctx.kit)
    shade(cell, paper)
    _line(cell, ctx, bb.SOCIAL_SAMPLE_HEADLINE, "h2", ink, True)
    _line(cell, ctx, bb.brand_name(ctx.kit), "small", ink, True)


def applications_part(ctx: _Ctx) -> None:
    from docx.shared import Mm

    from modules.documents.brand_board_miniatures import starter_miniatures

    _label(ctx, "Applications")
    drawn = starter_miniatures(ctx.kit)
    table = _table(ctx, 2, len(bb.APPLICATIONS))
    for column, (key, label) in enumerate(bb.APPLICATIONS):
        frame = table.cell(0, column)
        if key == bb.APPLICATION_SOCIAL:
            _social_cell(frame, ctx)
        else:
            _picture(frame, ctx, bb.miniature(drawn, key), width=Mm(DOCX_APP_MM))
        _line(table.cell(1, column), ctx, label)


PART_DOCX: Dict[str, Callable[[_Ctx], None]] = {
    "logo": logo_part,
    "variants": variants_part,
    "colours": colours_part,
    "type": type_part,
    "spacing": spacing_part,
    "voice": voice_part,
    "applications": applications_part,
}


def add_brand_part(doc: Any, block: Any, kit: Mapping[str, Any], font: Optional[str], load_image: ImageLoader) -> None:
    """Write one ``brand`` block into ``doc``, drawn from ``kit`` (the render-ready brand kit)."""
    PART_DOCX[block.part](_Ctx(doc, kit or {}, font, load_image))


__all__ = ["PART_DOCX", "add_brand_part"]
