"""A Word file's page, styles, letterhead and footer, from the PDF's design system (F356, 5 Oct).

F356: the owner asked for the Word files to match the PDFs. A block document's
.docx was python-docx's default template: Calibri headings in Word's theme blue,
no letterhead, no footer and no page numbers. Now, from ``design_tokens``:

* the page is A4 with the PDF's margins;
* Normal and Heading 1-3 take the type scale, the kit's font and the colour
  roles (title in the primary, section headings in the accent, body in the text
  colour), with Word's theme fonts and colours cleared so the kit's win;
* the letterhead (``letterhead_run``) is the first page's header: the logo on the
  left, the company block set right, an accent rule under both;
* every page's footer reads the company, the title and "Page X of Y", the page
  numbers as Word fields (PAGE, NUMPAGES), so they stay right when the file is edited.
"""
from __future__ import annotations

from typing import Any, Callable, Mapping, Optional

from . import design_tokens as t

PAGE_WIDTH_MM, PAGE_HEIGHT_MM = 210, 297
MARGIN_MM = {"top_margin": 20, "right_margin": 20, "bottom_margin": 24, "left_margin": 20}
HEADING_STYLES = {1: ("Heading 1", t.TITLE_PT), 2: ("Heading 2", t.H2_PT), 3: ("Heading 3", t.H3_PT)}
THEME_FONT_ATTRS = ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme")
RULE_EIGHTHS = 8  # a 1 pt rule, in Word's eighths of a point
LOGO_COLUMN_MM = 30
# What follows a paragraph border (w:pBdr) in a paragraph's properties.
PBDR_SUCCESSORS = (
    "w:shd", "w:tabs", "w:suppressAutoHyphens", "w:kinsoku", "w:wordWrap", "w:overflowPunct", "w:topLinePunct",
    "w:autoSpaceDE", "w:autoSpaceDN", "w:bidi", "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
    "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc", "w:textDirection", "w:textAlignment",
    "w:textboxTightWrap", "w:outlineLvl", "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange",
)


def rgb(colour: str):
    from docx.shared import RGBColor

    return RGBColor.from_string(colour.lstrip("#").upper())


def _clear_theme(style: Any) -> None:
    """Drop Word's theme font and colour from ``style``, so the kit's apply."""
    from docx.oxml.ns import qn

    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    for attr in THEME_FONT_ATTRS if fonts is not None else ():
        fonts.attrib.pop(qn(attr), None)
    colour = rpr.find(qn("w:color"))
    if colour is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            colour.attrib.pop(qn(attr), None)


def _set_font(style: Any, font: Optional[str], size: float, colour: str, bold: bool) -> None:
    from docx.shared import Pt

    _clear_theme(style)
    if font:
        style.font.name = font
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = rgb(colour)


def apply_styles(document: Any, kit: Mapping[str, Any], font: Optional[str]) -> None:
    """The page size and margins, and Normal and Heading 1-3 in the design system's type and colours."""
    from docx.shared import Mm, Pt

    roles = t.palette(kit)
    for section in document.sections:
        section.page_width, section.page_height = Mm(PAGE_WIDTH_MM), Mm(PAGE_HEIGHT_MM)
        for side, mm in MARGIN_MM.items():
            setattr(section, side, Mm(mm))
    normal = document.styles["Normal"]
    _set_font(normal, font, t.BODY_PT, roles.text, False)
    normal.paragraph_format.space_after = Pt(t.SPACE_2)
    normal.paragraph_format.line_spacing = t.LEADING - 0.25  # Word's single is already about 1.2x
    for level, (name, size) in HEADING_STYLES.items():
        style = document.styles[name]
        _set_font(style, font, size, roles.title if level == 1 else roles.heading, True)
        style.paragraph_format.space_before = Pt(t.SPACE_2 if level == 1 else t.SPACE_4)
        style.paragraph_format.space_after = Pt(t.SPACE_1 if level == 1 else t.SPACE_2)


EDGE_ORDER = ("top", "left", "bottom", "right")  # the order Word's schema wants them in


def cell_edges(cell: Any, edges: Mapping[str, tuple]) -> None:
    """Borders on a table cell: ``{"bottom": (colour, eighths of a point)}``, written in schema order.
    Call it before shading or vertical alignment, which come after the borders in a cell's properties."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for name in (edge for edge in EDGE_ORDER if edge in edges):
        colour, eighths = edges[name]
        edge = OxmlElement(f"w:{name}")
        for key, value in (("w:val", "single" if eighths else "nil"), ("w:sz", str(eighths)),
                           ("w:color", colour.lstrip("#"))):
            edge.set(qn(key), value)
        borders.append(edge)
    tc_pr.append(borders)


def _field(paragraph: Any, instruction: str, colour: str) -> Any:
    """A Word field (PAGE, NUMPAGES) as a run Word fills in when it lays the page out, in the caption's size."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), instruction)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    for tag, value in (("w:color", colour.lstrip("#")), ("w:sz", str(int(t.CAPTION_PT * 2)))):
        prop = OxmlElement(tag)
        prop.set(qn("w:val"), value)
        props.append(prop)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = "1"
    run.append(text)
    field.append(run)
    paragraph._p.append(field)
    return field


def _footer_paragraph(footer: Any, company: str, title: str, muted: str) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Mm, Pt

    paragraph = footer.paragraphs[0]
    width = Mm(PAGE_WIDTH_MM - MARGIN_MM["left_margin"] - MARGIN_MM["right_margin"])
    stops = paragraph.paragraph_format.tab_stops
    stops.add_tab_stop(width // 2, WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
    paragraph.add_run(f"{company}\t{title}\tPage ")
    _field(paragraph, "PAGE", muted)
    paragraph.add_run(" of ")
    _field(paragraph, "NUMPAGES", muted)
    for run in paragraph.runs:
        run.font.size = Pt(t.CAPTION_PT)
        run.font.color.rgb = rgb(muted)


def add_footer(document: Any, kit: Mapping[str, Any], company: str, title: str) -> None:
    """Every page's footer: the company, the title and "Page X of Y" (the first page's too)."""
    roles = t.palette(kit)
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer):
            _footer_paragraph(footer, company, title, roles.muted)


def full_width(table: Any, width: Optional[int] = None) -> None:
    """The table as wide as the text column, in twentieths of a point (not "auto", which some readers shrink)."""
    from docx.oxml.ns import qn
    from docx.shared import Mm

    width = width or Mm(PAGE_WIDTH_MM - MARGIN_MM["left_margin"] - MARGIN_MM["right_margin"])
    tbl_w = table._tbl.tblPr.find(qn("w:tblW"))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(int(width.pt * 20)))


def _rule_under(paragraph: Any, colour: str) -> None:
    """``paragraph`` as a thin line with an accent rule under it (the letterhead's rule)."""
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("w:val", "single"), ("w:sz", str(RULE_EIGHTHS)), ("w:space", "1"), ("w:color", colour.lstrip("#"))):
        bottom.set(qn(key), value)
    borders.append(bottom)
    p_pr.insert_element_before(borders, *PBDR_SUCCESSORS)  # Word's schema order inside a paragraph's properties
    paragraph.paragraph_format.space_after = Pt(t.SPACE_3)
    paragraph.add_run().font.size = Pt(1)


def add_letterhead(document: Any, kit: Mapping[str, Any], logo: Callable[[Any], None],
                   company: Callable[[Any], None]) -> None:
    """The first page's header: a two-cell row, ``logo`` writing into the left cell and
    ``company`` into the right, an accent rule under both."""
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Mm

    roles = t.palette(kit)
    section = document.sections[0]
    section.different_first_page_header_footer = True
    header = section.first_page_header
    width = Mm(PAGE_WIDTH_MM - MARGIN_MM["left_margin"] - MARGIN_MM["right_margin"])
    table = header.add_table(rows=1, cols=2, width=width)
    table.autofit = False
    full_width(table, width)
    left, right = table.rows[0].cells
    for column, cell, size in ((table.columns[0], left, Mm(LOGO_COLUMN_MM)), (table.columns[1], right, width - Mm(LOGO_COLUMN_MM))):
        column.width = cell.width = size
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    logo(left)
    company(right)
    for paragraph in right.paragraphs:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    rule = header.paragraphs[0]
    rule._p.addprevious(table._tbl)  # the row first; the header's own paragraph under it carries the rule
    _rule_under(rule, roles.rule)


__all__ = ["add_footer", "add_letterhead", "apply_styles", "cell_edges", "full_width", "rgb"]
